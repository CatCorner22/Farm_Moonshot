#!/usr/bin/env python3
"""Build Dora Ella Newnam Stone biography Word (.docx) from the master markdown.

Writes the authoritative file at
research/dora-ella-newnam/book/Dora-Ella-Newnam-Stone-A-Life.docx
and byte-identical working copies under sources/authored/.
"""

from __future__ import annotations

import hashlib
import re
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "research" / "dora-ella-newnam"
NARRATIVE = CASE / "book" / "Dora-Ella-Newnam-Stone-A-Life.md"
DOCX_PATH = NARRATIVE.with_suffix(".docx")
AUTHORED = CASE / "sources" / "authored"
LATEST = AUTHORED / "Dora-Ella-Newnam-Stone-A-Life-latest.docx"

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_ITALIC_SPLIT = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{3,}")


def _set_run_font(run, name: str = "Liberation Serif", size: Pt | None = None, italic: bool = False, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size
    run.italic = italic
    run.bold = bold


def add_rich_runs(paragraph, text: str, base_size: Pt | None = None) -> None:
    """Render **bold**, *italic*, `code`, and [label](url) into a paragraph."""
    text = LINK_RE.sub(r"\1", text)
    parts = BOLD_ITALIC_SPLIT.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            _set_run_font(run, size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, name="Liberation Mono", size=Pt(9))
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, size=base_size, italic=True)
        else:
            run = paragraph.add_run(part)
            _set_run_font(run, size=base_size)


def add_rich_paragraph(doc: Document, text: str, style: str | None = None, *, center: bool = False, italic_all: bool = False) -> None:
    p = doc.add_paragraph(style=style)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if italic_all:
        run = p.add_run(LINK_RE.sub(r"\1", text.strip("*").strip()))
        _set_run_font(run, italic=True)
        return
    add_rich_runs(p, text)


def is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped.startswith("|"):
        return False
    cells = [c.strip() for c in stripped.strip("|").split("|")]
    return all(TABLE_SEP_RE.match(c) or set(c) <= set("-: ") for c in cells if c)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_rich_runs(p, row[j] if j < len(row) else "", base_size=Pt(9))
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Serif"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x1C, 0x19, 0x17)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Serif")
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Liberation Serif"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x7C, 0x2D, 0x12)
        style.font.bold = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("Dora Ella Newnam Stone — a documented life, 1871–1926")
    _set_run_font(run, size=Pt(9), italic=True)
    run.font.color.rgb = RGBColor(0x57, 0x53, 0x4E)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Living descendants withheld. Confidence labels and claim IDs are part of the text.  ·  ")
    _set_run_font(run, size=Pt(8), italic=True)
    # PAGE field
    from docx.oxml import OxmlElement

    def add_page_field(paragraph) -> None:
        run1 = paragraph.add_run()
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        run1._r.append(fld1)
        run2 = paragraph.add_run()
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        run2._r.append(instr)
        run3 = paragraph.add_run()
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        run3._r.append(fld2)
        for r in (run1, run2, run3):
            _set_run_font(r, size=Pt(8))

    add_page_field(fp)

    core = doc.core_properties
    core.title = "Dora Ella Newnam Stone: A Documented Life, 1871–1926"
    core.subject = "Genealogical biography (GPS-cited). Living people redacted."
    core.author = "Farm_Moonshot case file — research/dora-ella-newnam"
    core.comments = (
        "Rebuilt from the markdown case book. Passes 62–88 filed (Pass 88 T623-840 crops not examined). "
        "Pope Roll 5 unread. Living people redacted."
    )


def build_docx() -> Path:
    if not NARRATIVE.is_file():
        raise SystemExit(f"missing narrative: {NARRATIVE}")
    lines = NARRATIVE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    i = 0
    pending_table: list[list[str]] | None = None

    def flush_table() -> None:
        nonlocal pending_table
        if pending_table:
            add_table(doc, pending_table)
            pending_table = None

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            flush_table()
            header = split_row(line)
            body: list[list[str]] = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_table_separator(lines[i]):
                    body.append(split_row(lines[i]))
                i += 1
            add_table(doc, [header] + body)
            continue

        if not stripped:
            flush_table()
            i += 1
            continue
        if stripped == "---":
            flush_table()
            doc.add_page_break()
            i += 1
            continue
        if stripped.startswith("# "):
            flush_table()
            p = doc.add_paragraph(style="Title")
            add_rich_runs(p, stripped[2:].strip(), base_size=Pt(22))
        elif stripped.startswith("## "):
            flush_table()
            doc.add_heading(LINK_RE.sub(r"\1", stripped[3:].strip()), level=1)
        elif stripped.startswith("### "):
            flush_table()
            doc.add_heading(LINK_RE.sub(r"\1", stripped[4:].strip()), level=2)
        elif stripped.startswith("- "):
            flush_table()
            add_rich_paragraph(doc, stripped[2:].strip(), style="List Bullet")
        elif stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**") and stripped.count("*") == 2:
            flush_table()
            add_rich_paragraph(doc, stripped, center=True, italic_all=True)
        else:
            flush_table()
            add_rich_paragraph(doc, stripped)
        i += 1

    flush_table()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = note.add_run(
        f"Master Word export {date.today().isoformat()}. "
        "Source: book/Dora-Ella-Newnam-Stone-A-Life.md. "
        "On-site films (Pope Roll 5; Danville Democrat 1926; 1885 marriage image) remain unread. "
        "Living people redacted."
    )
    _set_run_font(run, size=Pt(9), italic=True)

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    AUTHORED.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    return DOCX_PATH


def copy_working_files(src: Path) -> tuple[Path, Path, str, int]:
    digest = hashlib.sha256(src.read_bytes()).hexdigest()
    size = src.stat().st_size
    LATEST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, LATEST)
    dated = AUTHORED / f"Dora-Ella-Newnam-Stone-A-Life-book-copy-{date.today().isoformat()}.docx"
    shutil.copy2(src, dated)
    return LATEST, dated, digest, size


if __name__ == "__main__":
    path = build_docx()
    latest, dated, digest, size = copy_working_files(path)
    print(f"wrote {path} ({size} bytes)")
    print(f"sha256 {digest}")
    print(f"copies {latest} {dated}")
